"""Excel uploads: reading and checking a file, comparing it with the table, replacing the rows."""

import datetime
import io

import pytest
import sqlalchemy as sa
from openpyxl import Workbook, load_workbook

from catalyst.db import load_schema
from catalyst.db.sync import apply_changes
from catalyst.importers import excel
from catalyst.importers.excel import UploadError

from .conftest import make_schema

WEST, EAST = "GB108050008080", "GB108050008100"
HEADINGS = ("catchment_id", "measure", "year", "status", "value")
SCHEMA = load_schema()
EA_SCORES = SCHEMA.metadata.tables["ea_scores"]


def xlsx(*rows) -> bytes:
    workbook = Workbook()
    for row in rows:
        workbook.active.append(list(row))
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def problems(table, data) -> str:
    with pytest.raises(UploadError) as error:
        excel.read_upload(table, data)
    return "\n".join(error.value.problems)


def add_catchments(engine, *ids):
    with engine.begin() as conn:
        for catchment_id in ids:
            conn.execute(
                sa.text("INSERT INTO catchments (id, name, waterbody_type) VALUES (:id, :id, 'River')"),
                {"id": catchment_id},
            )


def ea_scores(engine) -> list[tuple]:
    with engine.connect() as conn:
        return [tuple(r) for r in conn.execute(sa.text("SELECT * FROM ea_scores ORDER BY catchment_id, measure, year"))]


# --- Reading a file ------------------------------------------------------------

def test_reads_rows_as_the_columns_types():
    rows = excel.read_upload(EA_SCORES, xlsx(
        ("value", "Catchment_ID ", "measure", "year", "status"),  # any order; case and spaces don't matter
        (0.16, WEST, " Fish ", 2025.0, "Poor"),
        (None, None, "", None, None),  # empty rows are skipped
        ("0.2", WEST, "Phosphate", "2022", None),
    ))
    assert [(r.line, r.values) for r in rows] == [
        (2, {"catchment_id": WEST, "measure": "Fish", "year": 2025, "status": "Poor", "value": 0.16}),
        (4, {"catchment_id": WEST, "measure": "Phosphate", "year": 2022, "status": None, "value": 0.2}),
    ]


def test_every_problem_is_reported_together():
    found = problems(EA_SCORES, xlsx(
        HEADINGS,
        (WEST, "Fish", "twenty", "Poor", 0.16),
        (WEST, None, 2025, "Poor", "lots"),
        (WEST, "Fish", 2022, "Good", 0.3),
        (WEST, "Fish", 2022, "Good", 0.4),
    ))
    for expected in [
        "Row 2, year: 'twenty' isn't a whole number.",
        "Row 3, measure: is empty, but needs a value.",
        "Row 3, value: 'lots' isn't a number.",
        f"Rows 4 and 5 are both for catchment_id '{WEST}', measure 'Fish', year 2022. Each combination can only appear once.",
    ]:
        assert expected in found


@pytest.mark.parametrize("headings, expected", [
    (("catchment_id", "measure", "year", "status"), "Missing column(s): value."),
    (HEADINGS + ("notes",), "Column(s) that aren't in the ea_scores table: notes."),
    (HEADINGS + ("year",), "The column year appears twice."),
])
def test_the_column_names_must_match_the_table(headings, expected):
    assert expected in problems(EA_SCORES, xlsx(headings, (WEST,)))


def test_an_empty_file_is_explained():
    assert "Row 1 of the first sheet should have the column names: catchment_id, measure" in problems(EA_SCORES, xlsx())


def test_a_file_that_isnt_excel_is_explained():
    assert "This isn't an Excel file (.xlsx)" in problems(EA_SCORES, b"catchment_id,measure\n")


def test_text_that_excel_turned_into_a_date_is_caught():
    found = problems(EA_SCORES, xlsx(HEADINGS, (WEST, datetime.datetime(2026, 1, 2), 2025, None, None)))
    assert "Row 2, measure: Excel has turned this into a date or time (2026-01-02)" in found


ALL_TYPES = make_schema("""
    tables:
      things:
        filled_by: excel
        columns:
          code: {type: "varchar(3)", primary_key: true}
          flag: {type: boolean}
          day: {type: date}
          small: {type: smallint}
""").metadata.tables["things"]


@pytest.mark.parametrize("cells, values", [
    (("A1", True, datetime.datetime(2025, 10, 1), 7), ("A1", True, datetime.date(2025, 10, 1), 7)),
    (("A1", "yes", "2025-10-01", "7"), ("A1", True, datetime.date(2025, 10, 1), 7)),
    (("A1", "No", "01/10/2025", 7.0), ("A1", False, datetime.date(2025, 10, 1), 7)),
    ((123, 0, None, None), ("123", False, None, None)),
])
def test_each_type_of_column(cells, values):
    (row,) = excel.read_upload(ALL_TYPES, xlsx(("code", "flag", "day", "small"), cells))
    assert tuple(row.values.values()) == values


@pytest.mark.parametrize("cells, expected", [
    (("ABCD", None, None, None), "code: is 4 characters long, but can be at most 3."),
    (("A1", "maybe", None, None), "flag: 'maybe' should be Yes or No."),
    (("A1", None, "next week", None), "day: 'next week' isn't a date."),
    (("A1", None, datetime.datetime(2025, 10, 1, 9, 30), None), "day: 2025-10-01 09:30 has a time as well as a date."),
    (("A1", None, None, 40000), "small: 40000 is too big for this column."),
    (("A1", None, None, 7.5), "small: 7.5 isn't a whole number."),
])
def test_values_that_dont_fit_the_column(cells, expected):
    assert f"Row 2, {expected}" in problems(ALL_TYPES, xlsx(("code", "flag", "day", "small"), cells))


def test_a_download_reads_back_exactly():
    rows = [
        {"catchment_id": WEST, "measure": "Fish", "year": 2025, "status": "Poor", "value": 0.16},
        {"catchment_id": WEST, "measure": "Phosphate", "year": 2025, "status": None, "value": 0.013},
    ]
    data = excel.write_table(EA_SCORES, rows)
    assert [r.values for r in excel.read_upload(EA_SCORES, data)] == rows

    sheet = load_workbook(io.BytesIO(data)).active
    assert [c.value for c in sheet[1]] == list(HEADINGS)
    assert sheet["A2"].number_format == "@"  # catchment IDs stay as text


# --- Comparing with the table and replacing it ------------------------------------

def test_an_upload_adds_changes_and_deletes_rows(engine):
    apply_changes(engine, SCHEMA)
    add_catchments(engine, WEST, EAST)
    with engine.begin() as conn:
        conn.execute(EA_SCORES.insert(), [
            {"catchment_id": WEST, "measure": "Fish", "year": 2025, "status": "Poor", "value": 0.16},
            {"catchment_id": WEST, "measure": "Fish", "year": 2022, "status": "Bad", "value": 0.1},
            {"catchment_id": EAST, "measure": "Fish", "year": 2025, "status": "Good", "value": 0.7},
        ])
    rows = excel.read_upload(EA_SCORES, xlsx(
        HEADINGS,
        (WEST, "Fish", 2025, "Poor", 0.16),      # unchanged
        (EAST, "Fish", 2025, "Moderate", 0.5),   # changed
        (EAST, "Fish", 2022, "Good", 0.65),      # added
    ))                                           # (WEST, Fish, 2022) is deleted

    with engine.begin() as conn:
        comparison = excel.compare(conn, EA_SCORES, rows)
        assert [r.line for r in comparison.added] == [4]
        assert [(c.row.line, c.columns) for c in comparison.changed] == [(3, ["status", "value"])]
        assert comparison.removed == [{"catchment_id": WEST, "measure": "Fish", "year": 2022, "status": "Bad", "value": 0.1}]
        assert comparison.unchanged == 1
        excel.replace(conn, EA_SCORES, comparison)

    assert ea_scores(engine) == [
        (WEST, "Fish", 2025, "Poor", 0.16),
        (EAST, "Fish", 2022, "Good", 0.65),
        (EAST, "Fish", 2025, "Moderate", 0.5),
    ]
    with engine.connect() as conn:
        assert not excel.compare(conn, EA_SCORES, rows).has_changes


def test_catchments_that_dont_exist_are_reported(engine):
    apply_changes(engine, SCHEMA)
    add_catchments(engine, WEST)
    rows = excel.read_upload(EA_SCORES, xlsx(
        HEADINGS,
        (WEST, "Fish", 2025, None, None),
        ("GB999", "Fish", 2025, None, None),
        ("GB999", "Fish", 2022, None, None),
    ))
    with engine.connect() as conn, pytest.raises(UploadError) as error:
        excel.compare(conn, EA_SCORES, rows)
    assert error.value.problems == ["catchment_id 'GB999' isn't in the catchments table (rows 3, 4)."]


def test_the_fingerprint_changes_when_the_table_does(engine):
    apply_changes(engine, SCHEMA)
    add_catchments(engine, WEST)
    rows = excel.read_upload(EA_SCORES, xlsx(HEADINGS, (WEST, "Fish", 2025, "Poor", 0.16)))
    with engine.begin() as conn:
        before = excel.compare(conn, EA_SCORES, rows).fingerprint
        assert excel.compare(conn, EA_SCORES, rows).fingerprint == before
        conn.execute(EA_SCORES.insert().values(catchment_id=WEST, measure="Fish", year=2025, status="Bad"))
        assert excel.compare(conn, EA_SCORES, rows).fingerprint != before
