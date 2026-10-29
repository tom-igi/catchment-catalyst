"""The Excel files: a table's download, and reading an edited copy back in.

The download has one sheet: the column names in row 1, then one row per row
of the table. An upload must look the same (the columns can be in any order),
and is read from its first sheet.
"""

from __future__ import annotations

import datetime
import io
import math
from collections.abc import Iterable, Mapping
from dataclasses import dataclass
from typing import Any

import sqlalchemy as sa
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

TRUE_WORDS = {"true", "yes", "y", "1"}
FALSE_WORDS = {"false", "no", "n", "0"}
INTEGER_LIMITS = {sa.SmallInteger: 2**15, sa.BigInteger: 2**63, sa.Integer: 2**31}


@dataclass
class FileRow:
    line: int  # row number in Excel; the column names are row 1
    values: dict[str, Any]  # column -> value


class UploadError(Exception):
    """The upload can't be used. `problems` lists everything wrong with it."""

    def __init__(self, problems: list[str]):
        self.problems = problems
        super().__init__("\n".join(problems))


def type_label(column: sa.Column) -> str:
    """What a column holds, in words, e.g. 'Whole number'."""
    t = column.type
    if isinstance(t, sa.Boolean):
        return "Yes or No"
    if isinstance(t, sa.Integer):
        return "Whole number"
    if isinstance(t, sa.Double):
        return "Number"
    if isinstance(t, sa.Date):
        return "Date"
    if isinstance(t, sa.String) and t.length:
        return f"Text (up to {t.length} characters)"
    return "Text"


# --- Writing the download ----------------------------------------------------

def write_table(table: sa.Table, rows: Iterable[Mapping[str, Any]]) -> bytes:
    """An Excel file of the table's rows."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = table.name[:31]  # Excel's limit
    columns = list(table.columns)
    sheet.append([c.name for c in columns])
    for row in rows:
        sheet.append([row[c.name] for c in columns])

    for i, column in enumerate(columns, start=1):
        letter = get_column_letter(i)
        # Applies to the whole column, including rows typed in later, so Excel
        # keeps text as typed (e.g. an ID stays an ID, not a number or date).
        number_format = "@" if isinstance(column.type, sa.String) else "yyyy-mm-dd" if isinstance(column.type, sa.Date) else None
        if number_format:
            sheet.column_dimensions[letter].number_format = number_format
            for (cell,) in sheet.iter_rows(min_row=2, min_col=i, max_col=i):
                cell.number_format = number_format
        widest = max((len(str(cell.value or "")) for (cell,) in sheet.iter_rows(min_col=i, max_col=i)), default=0)
        sheet.column_dimensions[letter].width = min(max(widest + 2, 10), 60)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"

    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


# --- Reading an upload -------------------------------------------------------

def read_upload(table: sa.Table, data: bytes) -> list[FileRow]:
    """Read and check an uploaded file for the table. Raises UploadError listing every problem."""
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # openpyxl raises many kinds of error for files it can't read
        raise UploadError([
            "This isn't an Excel file (.xlsx). If it's an older .xls file or a CSV file, "
            "open it in Excel and save it as an Excel Workbook (.xlsx)."
        ]) from None
    try:
        sheet = workbook.worksheets[0]
        sheet.reset_dimensions()  # read every row, even if the file says it has fewer
        lines = sheet.iter_rows(values_only=True)
        positions = _read_headings(table, next(lines, ()))
        problems: list[str] = []
        rows = []
        for line, cells in enumerate(lines, start=2):
            row = _read_row(table, positions, line, cells, problems)
            if row is not None:
                rows.append(row)
    finally:
        workbook.close()

    problems.extend(_duplicates(table, rows))
    if problems:
        raise UploadError(problems)
    return rows


def _read_headings(table: sa.Table, cells: tuple) -> dict[str, int]:
    """Where each of the table's columns is in the file."""
    names = [str(cell).strip().lower() if cell is not None else "" for cell in cells]
    expected = list(table.columns.keys())
    if not any(names):
        raise UploadError([f"Row 1 of the first sheet should have the column names: {', '.join(expected)}."])

    problems = []
    positions: dict[str, int] = {}
    for i, name in enumerate(names):
        if name in positions:
            problems.append(f"The column {name} appears twice.")
        elif name:
            positions[name] = i
    missing = [name for name in expected if name not in positions]
    unknown = [name for name in positions if name not in expected]
    if missing:
        problems.append(f"Missing column(s): {', '.join(missing)}.")
    if unknown:
        problems.append(f"Column(s) that aren't in the {table.name} table: {', '.join(unknown)}.")
    if problems:
        problems.append(f"Row 1 of the first sheet should have exactly these column names: {', '.join(expected)}.")
        raise UploadError(problems)
    return positions


def _read_row(table: sa.Table, positions: dict[str, int], line: int, cells: tuple, problems: list[str]) -> FileRow | None:
    """One row of the file, converted to the columns' types. None if the row is empty."""
    raw = {name: cells[i] if i < len(cells) else None for name, i in positions.items()}
    if all(value is None or (isinstance(value, str) and not value.strip()) for value in raw.values()):
        return None
    values = {}
    for column in table.columns:
        try:
            value = _convert(column, raw[column.name])
        except ValueError as e:
            problems.append(f"Row {line}, {column.name}: {e}")
            continue
        if value is None and not column.nullable:
            problems.append(f"Row {line}, {column.name}: is empty, but needs a value.")
        values[column.name] = value
    return FileRow(line, values)


def _duplicates(table: sa.Table, rows: list[FileRow]) -> list[str]:
    """Rows that have the same value(s) in the primary key or a column that must be unique."""
    problems = []
    for key in _unique_keys(table):
        first_seen: dict[tuple, int] = {}
        for row in rows:
            value = tuple(row.values.get(c) for c in key)
            if None in value:
                continue
            if value in first_seen:
                described = ", ".join(f"{c} {show_value(v)}" for c, v in zip(key, value))
                problems.append(
                    f"Rows {first_seen[value]} and {row.line} are both for {described}. "
                    f"Each {'combination' if len(key) > 1 else 'value'} can only appear once."
                )
            else:
                first_seen[value] = row.line
    return problems


def _unique_keys(table: sa.Table) -> list[tuple[str, ...]]:
    keys = [tuple(table.primary_key.columns.keys())]
    keys += [tuple(c.columns.keys()) for c in table.constraints if isinstance(c, sa.UniqueConstraint)]
    keys += [(c.name,) for c in table.columns if c.unique]
    keys += [tuple(i.columns.keys()) for i in table.indexes if i.unique]
    return list(dict.fromkeys(k for k in keys if k))


# --- Converting cells ----------------------------------------------------------

def _convert(column: sa.Column, value: Any) -> Any:
    """The cell's value as the column's type. Raises ValueError saying what's wrong."""
    if isinstance(value, str):
        value = value.strip()
        if not value:
            return None
    if value is None:
        return None
    t = column.type
    if isinstance(t, sa.Boolean):
        return _yes_no(value)
    if isinstance(t, sa.Integer):
        return _whole_number(value, t)
    if isinstance(t, sa.Double):
        return _number(value)
    if isinstance(t, sa.Date):
        return _date(value)
    if isinstance(t, sa.String):
        return _text(value, t.length)
    raise TypeError(f"Excel uploads can't fill {column.name}")  # schema.py stops this happening


def _text(value: Any, length: int | None) -> str:
    if isinstance(value, (datetime.date, datetime.time)):
        raise ValueError(
            f"Excel has turned this into a date or time ({show_value(value)}). "
            "Format the column as Text, then type the value again."
        )
    if isinstance(value, bool):
        text = "TRUE" if value else "FALSE"
    elif isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value)
    if length and len(text) > length:
        raise ValueError(f"is {len(text)} characters long, but can be at most {length}.")
    return text


def _whole_number(value: Any, t: sa.Integer) -> int:
    number = None
    if isinstance(value, int) and not isinstance(value, bool):
        number = value
    elif isinstance(value, float) and value.is_integer():
        number = int(value)
    elif isinstance(value, str):
        try:
            number = int(value)
        except ValueError:
            pass
    if number is None:
        raise ValueError(f"{show_value(value)} isn't a whole number.")
    limit = next(limit for kind, limit in INTEGER_LIMITS.items() if isinstance(t, kind))
    if not -limit <= number < limit:
        raise ValueError(f"{number} is too big for this column.")
    return number


def _number(value: Any) -> float:
    number = None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        number = float(value)
    elif isinstance(value, str):
        try:
            number = float(value)
        except ValueError:
            pass
    if number is None or not math.isfinite(number):
        raise ValueError(f"{show_value(value)} isn't a number.")
    return number


def _date(value: Any) -> datetime.date:
    if isinstance(value, datetime.datetime):
        if value.time() != datetime.time():
            raise ValueError(f"{show_value(value)} has a time as well as a date. Only the date is wanted.")
        return value.date()
    if isinstance(value, datetime.date):
        return value
    if isinstance(value, str):
        for parse in (datetime.date.fromisoformat, lambda v: datetime.datetime.strptime(v, "%d/%m/%Y").date()):
            try:
                return parse(value)
            except ValueError:
                pass
    raise ValueError(f"{show_value(value)} isn't a date. Use the form 2025-10-01 or 01/10/2025.")


def _yes_no(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    word = str(int(value) if isinstance(value, float) and value.is_integer() else value).lower()
    if word in TRUE_WORDS:
        return True
    if word in FALSE_WORDS:
        return False
    raise ValueError(f"{show_value(value)} should be Yes or No.")


def show_value(value: Any) -> str:
    if isinstance(value, datetime.datetime):
        return value.strftime("%Y-%m-%d %H:%M") if value.time() != datetime.time() else value.strftime("%Y-%m-%d")
    if isinstance(value, (datetime.date, datetime.time)):
        return value.isoformat()
    if isinstance(value, str):
        return f"'{value}'"
    return str(value)
